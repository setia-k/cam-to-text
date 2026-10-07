"""
Generates voucher_template.xlsx (needs `pip install openpyxl`, only for this script).

Sheet "Accumulative": columns A/B, rows unlimited (rules cover A1:B20000).
Sheet "Batch100":     rows 1-100, same A/B check, plus a marker in column C on
                      rows 25/50/75/100 (alternating blue / yellow).

Colors, in priority order:
  grey  - a cell in A or B that is not exactly <digits> digits (checked first,
          so a wrong-length B shows grey, not red/green)
  green - B matches A
  red   - B is empty or differs from A (the default)

<digits> lives in Accumulative!E1 (named range VoucherDigits, default 14);
change that one cell to change the rule on both sheets.

Columns A/B are Text-formatted so long numbers don't turn into scientific
notation. No header row, so row n is card n.
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

LAST_ROW_ACCUMULATIVE = 20000
DEFAULT_DIGITS = 14

RED = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
GREEN = PatternFill(start_color="99DD99", end_color="99DD99", fill_type="solid")
GREY = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")
MARK_COLORS = [PatternFill(start_color="9DC3E6", end_color="9DC3E6", fill_type="solid"),   # blue
               PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")]   # yellow

# Not exactly VoucherDigits characters, or any non-digit among them. Applied to
# A1:B<n> so the same relative formula checks both columns; empty cells pass.
BAD_LENGTH = ('AND(A1<>"",NOT(AND(LEN(A1)=VoucherDigits,'
              'SUMPRODUCT(--ISNUMBER(--MID(A1,ROW(INDIRECT("1:"&VoucherDigits)),1)))=VoucherDigits)))')


def setup_sheet(ws, last_row):
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    for col in ("A", "B"):
        for row in range(1, last_row + 1):
            ws[f"{col}{row}"].number_format = "@"

    ws.conditional_formatting.add(
        f"A1:B{last_row}", FormulaRule(formula=[BAD_LENGTH], fill=GREY, stopIfTrue=True))
    rng = f"B1:B{last_row}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=['AND(B1<>"",B1=A1)'], fill=GREEN, stopIfTrue=True))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=['TRUE'], fill=RED))


wb = Workbook()

acc = wb.active
acc.title = "Accumulative"
setup_sheet(acc, LAST_ROW_ACCUMULATIVE)
acc["D1"] = "Digits per voucher:"
acc["D1"].font = Font(bold=True)
acc["E1"] = DEFAULT_DIGITS
acc.column_dimensions["D"].width = 20
wb.defined_names["VoucherDigits"] = DefinedName("VoucherDigits", attr_text="Accumulative!$E$1")

batch = wb.create_sheet("Batch100")
setup_sheet(batch, 100)
batch.column_dimensions["C"].width = 6
for i, row in enumerate(range(25, 101, 25)):
    batch[f"C{row}"].fill = MARK_COLORS[i % 2]
batch["E1"] = "Digits (set on Accumulative!E1):"
batch["F1"] = "=VoucherDigits"
batch.column_dimensions["E"].width = 32

wb.save("voucher_template.xlsx")
print("Saved voucher_template.xlsx")
